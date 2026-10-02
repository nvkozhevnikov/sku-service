"""Offline verification of frozen XML/CSV/XLSX and correction evidence."""
import csv
import hashlib
import json
from pathlib import Path
import sys
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

def sha(path):
    with path.open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()

def verify(root):
    freeze=json.loads((root/'FREEZE_MANIFEST.json').read_text(encoding='utf-8'))
    xml=ET.parse(root/'UNIVERSAL_SUPPLIER_RC_FINAL.xml').getroot()
    existing=xml.findall('EXISTING/Product');new=xml.findall('NEW/Product');unresolved=xml.findall('UNRESOLVED/Product')
    assert (len(existing),len(new),len(unresolved))==(506,2,3899)
    assert len({n.get('sterbrust_product_id') for n in existing})==458
    assert all(n.get('sterbrust_product_id') is None and n.get('actionable')=='false' for n in new+unresolved)
    keys=[(n.get('namespace'),n.get('external_id')) for n in existing+new+unresolved]
    assert len(keys)==len(set(keys))==4407
    groups={n.get('new_group_id') for n in new}
    assert groups=={'NEWG-EE5998DC31FFB11ABF9C75DC','NEWG-892CFFBB6E0CCF2A054922B9'}
    for n in existing:
        commercial=json.loads(n.find('CommercialObservation').text)
        if commercial['price_state']=='price_on_request':assert commercial['numeric_price'] is None
    cards=json.loads((root/'PRODUCT_CARD_PAYLOADS.json').read_text(encoding='utf-8'))
    htz=next(p for p in cards if p['fields']['brand']['value']=='MEBOR')
    assert htz['fields']['model_execution']['value']=='HTZ-800 RS'
    base=htz['fields']['characteristics']['value']
    assert len(base)==7 and len(htz['optional_characteristics'])==4
    assert all(p['configuration']=='BASE_TABLE' for p in base)
    assert {p['value'] for p in base if p['normalized_property']=='pulley_width'}=={80.0}
    assert {p['value'] for p in base if p['normalized_property']=='blade_width'}=={100.0}
    uyym=next(p for p in cards if p['fields']['model_execution']['value']=='UYYM-120')
    weight=next(p for p in uyym['fields']['characteristics']['value'] if p['normalized_property'].startswith('weight'))
    assert 'unspecified' in weight['role']
    assert '146998' in (root/'CANONICAL_DUPLICATE_CORRECTION_PACK.json').read_text()
    assert '147023' in (root/'CANONICAL_DUPLICATE_CORRECTION_PACK.json').read_text()
    assert json.loads((root/'CANONICAL_DATA_ERROR_PACK.json').read_text(encoding='utf-8'))
    csv_counts={}
    for file in root.glob('*.csv'):
        with file.open(encoding='utf-8',newline='') as f:csv_counts[file.name]=sum(1 for _ in csv.DictReader(f))
    assert csv_counts=={'EXISTING.csv':506,'OFFERS.csv':506,'NEW_FULL_CARDS.csv':2,'REVIEW.csv':3640,'CONFLICT.csv':259,'SOURCE_TRACE.csv':4407}
    wb=load_workbook(root/'UNIVERSAL_SUPPLIER_RC_QA.xlsx',read_only=True,data_only=True)
    assert wb.sheetnames==['Сводка','Existing','Offers','NEW full cards','Review','Conflict','Source trace']
    assert [wb['Сводка'].cell(i,2).value for i in range(5,11)]==[506,2,3640,259,4407,458]
    for name,count in [('Existing',506),('Offers',506),('NEW full cards',2),('Review',3640),('Conflict',259),('Source trace',4407)]:
        # Artifact Tool omits optional XLSX dimension metadata. Count real rows,
        # never equate an absent max_row with missing source data.
        refs=[r[0] for r in wb[name].iter_rows(min_row=6,values_only=True) if r[0] is not None]
        assert len(refs)==len(set(refs))==count
    wb.close()
    for name,expected in freeze['input_sha256'].items():assert sha(root/'accepted'/name)==expected
    manifest={p.name:{'sha256':sha(p),'bytes':p.stat().st_size} for p in sorted(root.iterdir())
              if p.is_file() and p.suffix in {'.xml','.csv','.xlsx','.json'} and p.name not in {'SHA256_MANIFEST.json','VERIFIED.json','WORKBOOK_TABLES.json'}}
    (root/'SHA256_MANIFEST.json').write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
    result={'status':'REPOSITORY_VERIFIED_FROZEN_EXPORT','summary':freeze['summary'],'csv_counts':csv_counts,
            'xml_sha256':sha(root/'UNIVERSAL_SUPPLIER_RC_FINAL.xml'),
            'xlsx_sha256':sha(root/'UNIVERSAL_SUPPLIER_RC_QA.xlsx'),
            'final_manifest_sha256':sha(root/'SHA256_MANIFEST.json'),
            'http':0,'sql_writes':0,'restore_performed':False,'excel_native_engine_tested':False}
    (root/'VERIFIED.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result,indent=2))
if __name__=='__main__':verify(Path(sys.argv[1]))
