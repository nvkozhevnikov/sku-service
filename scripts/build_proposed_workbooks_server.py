"""Portable server fallback for diagnostic proposal workbooks.

The desktop authoring runtime is not installed in the server image. This
builder uses the already declared openpyxl dependency and accepts only the
validated, non-ESOL proposal JSON produced by generate_proposed_supplier_files.
"""

from __future__ import annotations

from datetime import datetime
import json
from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SPECS = (
    ("STERBRUST_UPDATE_EXISTING.xlsx", "Существующие товары: предложение обновления", "Existing", "existing", (
        ("Sterbrust ID", "sterbrust_product_id"), ("Модель", "model"),
        ("Название Sterbrust", "sterbrust_name"), ("Цена RUB", "price"),
        ("Валюта", "currency"), ("Наличие", "availability"),
        ("Предложенный источник", "selected_source_proposed"),
        ("ID у поставщика", "source_external_id"), ("Наблюдение UTC", "observed_at"),
        ("Статус", "selection_status"), ("URL карточки", "source_url"))),
    ("STERBRUST_CREATE_NEW_CANDIDATES.xlsx", "Новые товары: подтверждённые кандидаты", "New", "new", (
        ("NEW_CANDIDATE_ID", "new_candidate_id"), ("Источник", "source"),
        ("ID у источника", "external_id"), ("Модель", "model"),
        ("Исполнение (полная модель)", "execution"), ("Предлагаемое название", "proposed_name"),
        ("SECTION_ID", "section_id"), ("Путь раздела", "section_path"),
        ("Причина отсутствия в каталоге", "absence_evidence"), ("URL карточки", "source_url"),
        ("Источников", "source_offers"))),
    ("STERBRUST_REVIEW.xlsx", "Товары, требующие проверки", "Review", "review", (
        ("Источник", "source"), ("ID у источника", "external_id"), ("Модель", "model"),
        ("Название", "name"), ("Статус", "export_status"), ("Причина", "export_reason"),
        ("Признаки конфликта", "match_conflicts"), ("Результат matching", "match_status"),
        ("Предлагаемый Sterbrust ID", "proposed_sterbrust_id"), ("Цена", "price"),
        ("Состояние цены", "price_state"), ("Валюта", "currency"),
        ("Наличие", "availability"), ("URL карточки", "source_url"))),
)


def _value(row: dict, key: str):
    value = row.get(key)
    if key == "source_offers":
        return len(value or [])
    if key == "selection_status" and value == "PROPOSED_ONLY_SUPPLIERS_DISABLED":
        return "Предложение, не применено"
    if key == "match_conflicts":
        return "; ".join(value) if isinstance(value, list) else value or ""
    if key in {"price", "section_id"}:
        return float(value) if value not in (None, "") else None
    if key == "observed_at" and value:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).replace(tzinfo=None)
    return "" if value is None else str(value)


def build(input_path: Path, output_dir: Path) -> dict[str, int]:
    payload = json.loads(input_path.read_text(encoding="utf-8"))
    if payload.get("diagnostic_only") is not True or payload.get("esol_payload") is not False:
        raise ValueError("proposal workbook source is not a diagnostic-only payload")
    tables = payload.get("tables") or {}
    if any(not isinstance(tables.get(key), list) for key in ("existing", "new", "review")):
        raise ValueError("proposal workbook tables are incomplete")
    output_dir.mkdir(parents=True, exist_ok=True)
    counts = {}
    for filename, title, sheet_name, table_key, columns in SPECS:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = "18324F"
        sheet["A2"] = title
        sheet["A2"].font = Font(name="Arial", size=14, bold=True, color="18324F")
        sheet["A3"] = "Диагностический файл. Не является ESOL import payload и не изменяет Sterbrust."
        sheet["A3"].font = Font(name="Arial", size=10, italic=True, color="5B6672")
        for column, (label, key) in enumerate(columns, 1):
            cell = sheet.cell(5, column, label)
            cell.fill = PatternFill("solid", fgColor="18324F")
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center")
            width = 64 if key == "source_url" else 50 if key in {"section_path", "export_reason", "match_conflicts"} else 48 if key in {"name", "sterbrust_name", "proposed_name"} else 26
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.row_dimensions[5].height = 27
        rows = tables[table_key]
        if not rows:
            sheet["A6"] = "Подтверждённых записей для этого файла нет."
        for row_index, row in enumerate(rows, 6):
            for column, (_, key) in enumerate(columns, 1):
                cell = sheet.cell(row_index, column, _value(row, key))
                cell.font = Font(name="Arial", size=10, color="17212B")
                cell.alignment = Alignment(vertical="top", wrap_text=key in {
                    "name", "sterbrust_name", "proposed_name", "export_reason", "match_conflicts"})
                if key == "price":
                    cell.number_format = "#,##0.00"
                elif key == "section_id":
                    cell.number_format = "0"
                elif key == "observed_at":
                    cell.number_format = "yyyy-mm-dd hh:mm"
            sheet.row_dimensions[row_index].height = 36 if table_key == "review" else 30
        sheet.freeze_panes = "A6"
        sheet.auto_filter.ref = f"A5:{get_column_letter(len(columns))}{max(6, 5 + len(rows))}"
        workbook.save(output_dir / filename)
        counts[filename] = len(rows)
    return counts


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_proposed_workbooks_server.py PROPOSAL_TABLES.json OUT_DIR")
    print(json.dumps(build(Path(sys.argv[1]), Path(sys.argv[2])), ensure_ascii=False))
