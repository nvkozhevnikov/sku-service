"""Security review 2026-10, group 1: formula injection, blocking handlers, javascript: links."""
from __future__ import annotations

import csv
import inspect
import json
from pathlib import Path

from fastapi.testclient import TestClient
from jinja2 import ChainableUndefined, Environment, FileSystemLoader
from openpyxl import load_workbook

from universal_supplier.control_plane.admin_store import InMemoryAdminStore
from universal_supplier.control_plane.queue import InMemoryJobQueue
from universal_supplier.control_plane.rate_control import DomainRatePolicy
from universal_supplier.control_plane.web import (
    INTERACTIVE_TOTAL_BUDGET_SECONDS, create_app, interactive_fetch_policy, safe_http_url,
)
from universal_supplier.models import FetchRecord
from tests.test_stage6d2 import RenderStore

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "universal_supplier/control_plane/templates"
HOSTILE = ("=HYPERLINK(\"http://evil.example\",\"x\")", "+cmd|' /C calc'!A0", "@SUM(1)", "-2+3")


# --- 1. Formula injection -------------------------------------------------

def test_server_workbook_string_cells_are_never_formulas(tmp_path: Path):
    from scripts.build_proposed_workbooks_server import build
    source = tmp_path / "proposal.json"
    source.write_text(json.dumps({
        "diagnostic_only": True, "esol_payload": False,
        "tables": {
            "existing": [{"sterbrust_product_id": 1, "model": HOSTILE[0], "sterbrust_name": HOSTILE[1],
                          "price": 10, "source_url": "https://ok.example/a"}],
            "new": [{"new_candidate_id": "NEWC-1", "model": HOSTILE[1], "proposed_name": HOSTILE[0],
                     "source_offers": []}],
            "review": [{"source": "x", "model": HOSTILE[0], "name": HOSTILE[1],
                        "match_conflicts": [HOSTILE[2]]}],
        },
    }), encoding="utf-8")
    build(source, tmp_path / "out")
    existing = load_workbook(tmp_path / "out" / "STERBRUST_UPDATE_EXISTING.xlsx").active
    assert existing["B6"].data_type == "s" and existing["B6"].value == "'" + HOSTILE[0]
    assert existing["C6"].data_type == "s" and existing["C6"].value == "'" + HOSTILE[1]
    assert existing["D6"].value == 10  # numbers stay numbers
    new = load_workbook(tmp_path / "out" / "STERBRUST_CREATE_NEW_CANDIDATES.xlsx").active
    assert new["D6"].value == "'" + HOSTILE[1] and new["F6"].value == "'" + HOSTILE[0]
    review = load_workbook(tmp_path / "out" / "STERBRUST_REVIEW.xlsx").active
    assert review["C6"].value == "'" + HOSTILE[0] and review["D6"].value == "'" + HOSTILE[1]
    assert review["G6"].value == "'" + HOSTILE[2]
    for sheet in (existing, new, review):
        assert all(cell.data_type != "f" for row in sheet.iter_rows() for cell in row)


def test_matching_report_csv_cells_are_neutralised():
    from scripts.generate_proposed_supplier_files import _csv_cell
    assert _csv_cell(HOSTILE[0]) == "'" + HOSTILE[0]
    assert _csv_cell(HOSTILE[1]) == "'" + HOSTILE[1]
    assert _csv_cell("BMS-230DG") == "BMS-230DG"
    assert _csv_cell(12.5) == 12.5 and _csv_cell(None) is None
    assert _csv_cell(["=1+1"]) == '["=1+1"]'  # JSON text starts with '[' and stays literal
    assert _csv_cell("a\x00b") == "ab"


def test_final_rc_csv_rows_are_neutralised(tmp_path: Path):
    from scripts.assemble_final_rc import csv_row
    row = csv_row({"model": HOSTILE[0], "name": HOSTILE[1], "price": 5, "id": None})
    path = tmp_path / "x.csv"
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(row)); writer.writeheader(); writer.writerow(row)
    parsed = next(csv.DictReader(path.open(encoding="utf-8")))
    assert parsed["model"].startswith("'=") and parsed["name"].startswith("'+")
    assert row["price"] == 5 and row["id"] is None


# --- 2. Blocking network handlers -----------------------------------------

def _route_endpoint(app, path: str, method: str = "POST"):
    return next(route.endpoint for route in app.routes
                if getattr(route, "path", None) == path and method in getattr(route, "methods", ()))


def test_network_handlers_are_not_coroutines():
    app = create_app(store=RenderStore(), queue=InMemoryJobQueue(), admin_store=InMemoryAdminStore(), auth_required=False)
    for path in ("/parser-inspector", "/products/{supplier_code}/{external_id}/reparse"):
        assert not inspect.iscoroutinefunction(_route_endpoint(app, path)), path


def test_interactive_fetch_policy_is_clamped():
    clamped = interactive_fetch_policy(DomainRatePolicy(max_attempts=5, request_timeout_seconds=40,
                                                        max_run_duration_seconds=7200))
    assert clamped.max_attempts == 2 and clamped.request_timeout_seconds == 15
    assert clamped.max_run_duration_seconds <= 60 and INTERACTIVE_TOTAL_BUDGET_SECONDS <= 60
    gentle = interactive_fetch_policy(DomainRatePolicy(max_attempts=1, request_timeout_seconds=5))
    assert gentle.max_attempts == 1 and gentle.request_timeout_seconds == 5


def test_parser_inspector_uses_clamped_transport_and_still_reads_form(monkeypatch):
    seen = {}

    class FakeTransport:
        def __init__(self, **kwargs): seen.update(kwargs)
        def get_validated(self, url, validator, **kwargs):
            raise RuntimeError("stop after transport construction")

    monkeypatch.setattr("universal_supplier.http.HttpTransport", FakeTransport)
    monkeypatch.setattr("universal_supplier.control_plane.web.validate_supplier_url", lambda url, allowed, **kw: url)
    store = RenderStore(); store.supplier["rate_settings"] = {"max_attempts": 4, "request_timeout_seconds": 40}
    app = create_app(store=store, queue=InMemoryJobQueue(), admin_store=InMemoryAdminStore(), auth_required=False)
    response = TestClient(app).post("/parser-inspector", data={"csrf_token": "test-csrf", "supplier": "optimum",
                                                               "url": "https://example.invalid/p"})
    assert response.status_code == 400 and "Проверка парсера завершилась ошибкой" in response.text
    assert seen["domain_policy"].max_attempts == 2 and seen["domain_policy"].request_timeout_seconds == 15
    assert seen["deadline"] is not None


# --- 3. javascript: links ---------------------------------------------------

def test_safe_http_url_filter():
    assert safe_http_url("https://partner-st.ru/a?b=1") == "https://partner-st.ru/a?b=1"
    assert safe_http_url(" http://x.example/ ") == "http://x.example/"
    for hostile in ("javascript:alert(1)", "JaVaScRiPt:alert(1)", " javascript:alert(1)", "data:text/html,x",
                    "vbscript:x", "//evil.example/x", "/relative", "https:///nohost", "java\nscript:alert(1)",
                    "https://ok.example/\x00", None, ""):
        assert safe_http_url(hostile) == "", hostile


class XssStore(RenderStore):
    def product(self, supplier_code, external_id):
        row = super().product(supplier_code, external_id)
        row.update(name="<script>alert(1)</script>", source_url="javascript:alert(1)", media=[
            {"source_url": "javascript:alert(1)", "title": "<script>alert(2)</script>"},
            {"source_url": "https://partner-st.ru/doc.pdf", "title": "Паспорт"}])
        return row


def test_product_detail_never_renders_javascript_href():
    app = create_app(store=XssStore(), queue=InMemoryJobQueue(), admin_store=InMemoryAdminStore(), auth_required=False)
    html = TestClient(app).get("/products/optimum/3").text
    assert 'href="javascript:' not in html.lower() and "<script>alert" not in html
    assert "&lt;script&gt;alert(2)&lt;/script&gt;" in html
    assert 'href="https://partner-st.ru/doc.pdf"' in html


def _render(template: str, **context) -> str:
    from universal_supplier.control_plane import web
    env = Environment(loader=FileSystemLoader(TEMPLATES), autoescape=True, undefined=ChainableUndefined)
    env.filters.update(ru=web.russian_label, product_kind_ru=web.product_kind_label,
                       relation_ru=web.relation_type_label, safe_http_url=web.safe_http_url)
    env.globals["url_for"] = lambda *args, **kwargs: "/static/x"
    return env.get_template(template).render(request=None, current_user=None, **context)


def test_operator_templates_drop_non_http_links():
    bad = "javascript:alert(1)"
    rendered = [
        _render("final_release.html", release={"rows": [{"source_url": bad, "name": "N"}], "summary": {}, "page": 1, "pages": 1}),
        _render("commercial_proposals.html", unavailable="", proposals={
            "counts": {}, "tables": {"existing": [{"source_url": bad}], "new": [{"source_url": bad}],
                                     "review": [{"source_url": bad}]}}),
        _render("commercial_qa_run.html", run={"rows": [{"status": "DRY_RUN", "url": bad}]}),
    ]
    for html in rendered:
        assert "href=\"javascript:" not in html.lower() and "href=\"\"" not in html
    good = _render("final_release.html", release={"rows": [{"source_url": "https://ok.example/p", "name": "N"}], "summary": {}, "page": 1, "pages": 1})
    assert 'href="https://ok.example/p"' in good


def test_partner_st_ingest_drops_non_http_media():
    from universal_supplier.adapters.partner_st import PartnerStAdapter
    html = """<html><head><script type="application/ld+json">{"@type":"Product","name":"P",
      "image":["javascript:alert(1)","/img/a.jpg"]}</script></head><body><h1>P</h1>
      <input name="product_data[5][product_id]" value="5">
      <a href="javascript:alert(1)//x.pdf">doc</a><a href="/files/manual.pdf">manual</a>
      <iframe src="javascript:alert(2)"></iframe><iframe src="https://video.example/v"></iframe></body></html>"""
    card = PartnerStAdapter().parse_product(FetchRecord("https://partner-st.ru/p", "https://partner-st.ru/p", 200, (), 1, 1, html))
    urls = [item.url for item in card.media]
    assert urls and all(url.startswith(("http://", "https://")) for url in urls)
    assert "https://partner-st.ru/img/a.jpg" in urls and "https://partner-st.ru/files/manual.pdf" in urls
    assert "https://video.example/v" in urls


def test_kami_ingest_drops_non_http_images():
    from universal_supplier.adapters.kami import KamiAdapter
    base = "https://www.stanki.ru/catalog/test/item/"
    source = (ROOT / "tests/fixtures/kami/ironmac_single_sale.html").read_text(encoding="utf-8")
    hostile = source.replace('productPreview__picture" href="/upload/', 'productPreview__picture" href="javascript:alert(1)//upload/', 1)
    assert hostile != source
    before = [m.url for m in KamiAdapter().parse_product(FetchRecord(base, base, 200, (), 1, 0, source)).media if m.media_type == "image"]
    after = [m.url for m in KamiAdapter().parse_product(FetchRecord(base, base, 200, (), 1, 0, hostile)).media if m.media_type == "image"]
    assert len(after) == len(before) - 1
    assert all(url.startswith("https://") for url in after)
