import json
from pathlib import Path
from zipfile import ZipFile

from scripts.build_proposed_workbooks_server import build


def test_server_workbook_builder_keeps_missing_price_distinct_from_zero(tmp_path: Path):
    source = tmp_path / "proposal.json"
    source.write_text(json.dumps({
        "diagnostic_only": True, "esol_payload": False,
        "tables": {
            "existing": [{"sterbrust_product_id": 123, "model": "BMS-230DG",
                          "price": None, "currency": None, "source_external_id": "19240"}],
            "new": [{"new_candidate_id": "NEWC-TEST", "source": "intervesp", "external_id": "9028",
                     "model": "BMSY-440DGH WP2", "execution": "BMSY-440DGH WP2",
                     "section_id": 42, "source_offers": []}],
            "review": [{"source": "beka_mak", "external_id": "x", "model": "BMSY-320DGH",
                        "price": None, "match_conflicts": ["visible model differs"]}],
        },
    }), encoding="utf-8")
    counts = build(source, tmp_path / "out")
    assert list(counts.values()) == [1, 1, 1]
    from openpyxl import load_workbook
    existing = load_workbook(tmp_path / "out" / "STERBRUST_UPDATE_EXISTING.xlsx", read_only=True).active
    assert existing["D6"].value is None
    assert existing["H6"].value == "19240"
    new = load_workbook(tmp_path / "out" / "STERBRUST_CREATE_NEW_CANDIDATES.xlsx", read_only=True).active
    assert new["A6"].value == "NEWC-TEST"
    assert new["B6"].value == "intervesp"
    assert new["C6"].value == "9028"
    review = load_workbook(tmp_path / "out" / "STERBRUST_REVIEW.xlsx", read_only=True).active
    assert review["G6"].value == "visible model differs"
    for filename in counts:
        with ZipFile(tmp_path / "out" / filename) as archive:
            assert archive.testzip() is None
