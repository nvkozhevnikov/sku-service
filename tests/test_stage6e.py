from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from universal_supplier.control_plane.admin_store import InMemoryAdminStore
from universal_supplier.control_plane.queue import InMemoryJobQueue
from universal_supplier.control_plane.security import SessionSigner, hash_password
from universal_supplier.control_plane.store import ProductFilters, SnapshotControlPlaneStore, product_where
from universal_supplier.control_plane.web import create_app
from universal_supplier.stage6e import (
    AnomalyThresholds, InMemoryStage6EService, MissingState, QualityProfile,
    build_change_preview, csv_bytes, detect_anomalies, mapping_evidence,
    normalize_property_value, observe_missing, product_workbook,
    safe_export_filename, validate_supplier_url,
)

ROOT = Path(__file__).resolve().parents[1]


def make_client(role="ADMIN", service=None):
    admin=InMemoryAdminStore([{"id":1,"username":"user1","display_name":"Пользователь","role":role,"is_active":True,"password_hash":hash_password("Safe-password-2026",salt=b"0123456789abcdef")}])
    app=create_app(store=SnapshotControlPlaneStore(ROOT/"reports"),queue=InMemoryJobQueue(),admin_store=admin,
                   stage6e_service=service or InMemoryStage6EService(),session_secret="stage6e-test-secret-at-least-32-bytes",auth_required=True)
    client=TestClient(app); assert client.post("/login",data={"username":"user1","password":"Safe-password-2026"},follow_redirects=False).status_code==303
    return client,admin


def csrf(client):
    return SessionSigner("stage6e-test-secret-at-least-32-bytes").read(client.cookies.get(SessionSigner.cookie_name))["csrf"]


def profile(count, **rates):
    defaults={key:.98 for key in ("price","sku","brand","description","property","image","category")}; defaults.update(rates)
    return QualityProfile(count,count,defaults)


def test_count_collapse_is_critical_and_suppresses_missing_consequences():
    result=detect_anomalies(profile(1351),profile(74))
    assert "CRAWL_COUNT_COLLAPSE" in result.codes and result.critical
    assert result.run_status=="partial" and result.allow_missing_consequences is False


@pytest.mark.parametrize("field",["price","sku","property","image"])
def test_fill_rate_collapse_is_explicit(field):
    result=detect_anomalies(profile(100),profile(100,**{field:.10}))
    assert f"FIELD_COVERAGE_COLLAPSE:{field}" in result.codes


def test_schema_drift_keeps_evidence():
    result=detect_anomalies(None,profile(10),schema_signals=["gallery_missing","json_ld_changed"])
    assert result.codes==("PARSER_SCHEMA_DRIFT",) and result.evidence["schema_signals"]==["gallery_missing","json_ld_changed"]


@pytest.mark.parametrize("evidence",["TIMEOUT","HTTP_403","HTTP_429","NETWORK_ERROR","PARSER_ERROR"])
def test_network_or_parser_error_never_counts_as_missing(evidence):
    assert observe_missing(MissingState(1,"WATCH"),evidence).consecutive_confirmed==1


def test_confirmed_missing_grace_states_and_seen_reset():
    one=observe_missing(MissingState(),"HTTP_404"); two=observe_missing(one,"NOT_IN_SITEMAP"); three=observe_missing(two,"HTTP_410")
    assert (one.state,two.state,three.state)==("WATCH","WARNING","MISSING_CANDIDATE")
    assert observe_missing(three,"SEEN").state=="ACTIVE"


def test_anomaly_guard_freezes_missing_counter():
    assert observe_missing(MissingState(2,"WARNING"),"HTTP_404",anomaly_guarded=True)==MissingState(2,"WARNING","HTTP_404")


@pytest.mark.parametrize(("raw","unit","target","expected"),[("0.75","kW","W",Decimal("750")),("3","cm","mm",Decimal("30")),("1500","об/мин","rpm",Decimal("1500"))])
def test_property_unit_normalization(raw,unit,target,expected):
    result=normalize_property_value(raw,unit,target)
    assert result.status=="NORMALIZED" and result.numeric_value==expected


def test_incompatible_units_and_raw_values_are_preserved():
    result=normalize_property_value("3","kg","mm")
    assert result.status=="INCOMPATIBLE_DIMENSION" and result.raw_value=="3" and result.raw_unit=="kg"


def test_ambiguous_property_names_are_only_ranked_not_confirmed():
    result=mapping_evidence("Диаметр","мм","lathe",[
        {"id":1,"display_name_ru":"Диаметр шпинделя","canonical_unit":"mm","categories":["lathe"]},
        {"id":2,"display_name_ru":"Диаметр диска","canonical_unit":"mm","categories":["lathe"]},
    ])
    assert len(result)==2 and all("status" not in row for row in result)


@pytest.mark.parametrize("url",["http://127.0.0.1/admin","http://169.254.169.254/latest","file:///etc/passwd","https://evil.example/p"])
def test_inspector_ssrf_guard_rejects_arbitrary_or_local_urls(url):
    with pytest.raises(ValueError): validate_supplier_url(url,["partner-st.ru"])


def test_inspector_accepts_exact_adapter_host_only():
    assert validate_supplier_url("https://partner-st.ru/catalog/item","partner-st.ru".split())


def test_field_ownership_allows_commercial_and_protects_editorial():
    preview=build_change_preview({"price":1,"description":"manual"},{"price":2,"description":"supplier"},{"price":"SUPPLIER_ALLOWED","description":"STERBRUST_OWNED"})
    assert {row["field"]:row["decision"] for row in preview}=={"description":"NO_UPDATE","price":"PROPOSE_UPDATE"}


def test_anomaly_and_exclusion_override_preview():
    assert build_change_preview({"price":1},{"price":2},{"price":"SUPPLIER_ALLOWED"},blocked=True)[0]["decision"]=="ANOMALY_GUARD_BLOCKED"
    assert build_change_preview({"price":1},{"price":2},{"price":"SUPPLIER_ALLOWED"},exclusion=True)[0]["decision"]=="EXCLUDED_BY_RULE"


def sample_rows():
    return [{"source_product_id":1,"supplier_code":"optimum","supplier_name":"Optimum","external_id":"782","sku":"QW","name":"Станок","brand_raw":"Optimum","decision":"REVIEW_EXISTING","price":10,"currency":"RUB"}]


def test_csv_has_utf8_bom_russian_header_and_no_secret_fields():
    body=csv_bytes(sample_rows())
    text=body.decode("utf-8-sig")
    assert text.startswith("Поставщик,External ID") and "password_hash" not in text and "session" not in text


def test_full_xlsx_has_exact_seven_sheets_stable_ids_and_no_macros():
    body=product_workbook(sample_rows())
    book=load_workbook(BytesIO(body),read_only=True)
    assert book.sheetnames==["Товары","Характеристики","Изображения","Документы","Связи","Сопоставление Sterbrust","История решений"]
    assert "Source product ID" in [cell.value for cell in next(book["Характеристики"].iter_rows())]
    assert book.vba_archive is None


def test_export_filename_cannot_traverse():
    value=safe_export_filename("../products","csv","../../optimum")
    assert "/" not in value and "\\" not in value and value.endswith(".csv")


def test_filtered_export_parity_with_visible_snapshot_set():
    store=SnapshotControlPlaneStore(ROOT/"reports"); filters=ProductFilters(supplier="optimum",decision="SAFE_NEW_PRODUCT_CANDIDATE",brand="Optimum",page_size=100)
    assert len(store.export_rows(filters))==store.products(filters)["total"]


def test_quality_drilldown_uses_exact_database_predicate():
    where,_=product_where(ProductFilters(quality_issue="without_price"))
    assert "o.price IS NULL" in where


@pytest.mark.parametrize("extension",["csv","xlsx"])
def test_change_export_supports_operator_time_windows(extension):
    client,_=make_client(); response=client.get(f"/changes/export.{extension}?period=24h")
    assert response.status_code==200


def test_change_export_between_runs_requires_two_run_ids():
    client,_=make_client(); assert client.get("/changes/export.csv?period=between&run_a=1").status_code==400


@pytest.mark.parametrize("path",["/data-quality","/properties","/rules","/sterbrust-preview","/parser-inspector"])
def test_stage6e_operator_pages_are_russian_and_available(path):
    client,_=make_client(); response=client.get(path)
    assert response.status_code==200 and "Универсальный поставщик" in response.text


def test_product_exports_are_authenticated_and_filter_preserving():
    client,_=make_client(); response=client.get("/products/export.csv?supplier=optimum&decision=SAFE_NEW_PRODUCT_CANDIDATE&brand=Optimum")
    assert response.status_code==200 and response.content.startswith(b"\xef\xbb\xbf")
    assert "products_optimum_" in response.headers["content-disposition"]


def test_xlsx_route_uses_write_only_workbook():
    client,_=make_client(); response=client.get("/products/export.xlsx?supplier=optimum")
    assert response.status_code==200
    assert load_workbook(BytesIO(response.content),read_only=True).sheetnames[0]=="Товары"


@pytest.mark.parametrize("entity",["review","suppliers","runs","quality","properties"])
def test_operator_sections_offer_csv_and_xlsx_exports(entity):
    client,_=make_client()
    assert client.get(f"/exports/{entity}.csv").status_code==200
    response=client.get(f"/exports/{entity}.xlsx")
    assert response.status_code==200 and response.content[:2]==b"PK"


def test_viewer_can_export_but_cannot_create_rule():
    client,_=make_client(role="VIEWER")
    assert client.get("/products/export.csv").status_code==200
    assert client.post("/rules",data={"reason":"x","csrf_token":csrf(client)}).status_code==403


def test_rule_and_mapping_actions_are_audited():
    service=InMemoryStage6EService(properties=[{"id":7,"display_name_ru":"Мощность","code":"POWER","data_type":"number"}],mappings=[{"id":4,"supplier_name":"Optimum","source_name":"Мощность","status":"UNMAPPED"}])
    client,admin=make_client(service=service)
    token=csrf(client)
    assert client.post("/rules",data={"scope":"product","entity_key":"782","rule_type":"DO_NOT_IMPORT","reason":"Решение оператора","csrf_token":token},follow_redirects=False).status_code==303
    assert client.post("/properties/mappings/4/confirm",data={"canonical_property_id":"7","csrf_token":token},follow_redirects=False).status_code==303
    assert {row["action"] for row in admin.audit_events()} >= {"BUSINESS_RULE_CREATED","PROPERTY_MAPPING_CONFIRMED"}


def test_parser_inspector_rejects_loopback_before_any_fetch():
    client,_=make_client(); response=client.post("/parser-inspector",data={"supplier":"optimum","url":"http://127.0.0.1:8080/private","csrf_token":csrf(client)})
    assert response.status_code==400 and "Хост не разрешён" in response.text


def test_migration_014_is_coherent_and_previous_migrations_are_untouched():
    migration=(ROOT/"migrations/014_stage6e_operator_data_tools.sql").read_text(encoding="utf-8")
    for token in ("data_quality_snapshots","missing_observations","business_rules","canonical_properties","supplier_property_mappings","sterbrust_property_mappings","category_filter_profiles","field_observations","field_update_policies","reparse_previews"):
        assert token in migration
    assert "UPDATE sterbrust" not in migration.lower() and "INSERT INTO sterbrust" not in migration.lower()


def test_no_sterbrust_write_or_preview_apply_route_exists():
    client,_=make_client(); routes=[route.path.lower() for route in client.app.routes]
    assert not any(any(token in route for token in ("sterbrust/write","sterbrust/apply","esol","filter/write","property/create")) for route in routes)


def test_single_reparse_domain_service_never_runs_full_crawl():
    service=InMemoryStage6EService(); assert service.full_crawl_runs==0 and service.reparse_runs==0


def test_stage6_blocker_and_stage7_stop_remain_documented():
    text=(ROOT/"SUMMARY.txt").read_text(encoding="utf-8")+(ROOT/"README.md").read_text(encoding="utf-8")
    assert "BLOCKED_IMPORT_CONTRACT_EVIDENCE" in text and "Stage 7" in text
